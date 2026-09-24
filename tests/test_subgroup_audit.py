"""Whole-workbook checks: raw Excel assignments and every subgroup's delivery paths."""

import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from test_handlers import FakeCallback, FakeMessage, Users, callbacks

import app.bot.handlers as handlers
import app.calendar.service as calendar
from app.schedule.parser import ExcelScheduleParser
from app.schedule.service import ScheduleService

FILES = tuple(sorted(Path("tests/fixtures").glob("*.xlsx")))
ROOM = re.compile(r"\((онлайн|\d+)\s*\[(\d+)\](?:\s*,\s*(\d+))?\)", re.I)


@pytest.mark.parametrize("path", FILES, ids=lambda p: p.stem)
def test_every_excel_assignment_is_preserved(path):
    # Independent oracle: read room/subgroup annotations directly, including merges.
    expected = set()
    book = load_workbook(path, data_only=True)
    for sheet in book:
        if not re.fullmatch(r"\d курс", sheet.title):
            continue
        values = {(c.row, c.column): c.value for row in sheet for c in row}
        for area in sheet.merged_cells.ranges:
            value = values[area.min_row, area.min_col]
            for row in range(area.min_row, area.max_row + 1):
                for column in range(area.min_col, area.max_col + 1):
                    values[row, column] = value
        groups = {}
        day = None
        for row in range(1, sheet.max_row + 1):
            for column in range(3, sheet.max_column + 1):
                value = str(values[row, column] or "").strip()
                if re.fullmatch(r"[А-ЯЁA-Z]+-\d{2}-\d+", value):
                    groups[column] = value
            match = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", str(values[row, 1]))
            if match:
                d, m, y = map(int, match.groups())
                day = date(y, m, d)
            pair_text = str(values[row, 2] or "").strip()
            if day is None or not re.match(r"\d+\s+.*", pair_text, re.S):
                continue
            pair = int(pair_text.split()[0])
            for column, group in groups.items():
                raw = str(values[row, column] or "").strip()
                marks = ROOM.findall(raw)
                if raw and not marks:
                    expected.add((group, day, pair, None, None))
                for room, building, number in marks:
                    location = "онлайн" if room.lower() == "онлайн" else f"{room}[{building}]"
                    expected.add((group, day, pair, location, int(number) if number else None))
    book.close()
    schedule = ExcelScheduleParser().parse(path)
    actual = {(x.group, x.date, x.pair_number, x.location, x.subgroup) for x in schedule.lessons}
    assert not expected - actual, ("missing", expected - actual)
    assert not actual - expected, ("unexpected", actual - expected)


@pytest.fixture(scope="module", params=FILES, ids=lambda p: p.stem)
def schedule(request):
    return ExcelScheduleParser().parse(request.param)


async def test_all_groups_subgroups_days_weeks_and_calendar(schedule, monkeypatch):
    monkeypatch.setattr(handlers, "Message", FakeMessage)
    timezone = ZoneInfo("Asia/Yekaterinburg")
    for course, groups in schedule.courses.items():
        for group in groups:
            lessons = schedule.for_group(group)
            numbers = sorted({x.subgroup for x in lessons if x.subgroup is not None})
            if numbers:
                users = Users()
                router = handlers.build_router(users, ScheduleService(schedule), timezone)
                registration = FakeCallback(f"group:{course}:{group}")
                await callbacks(router, "callback_query")["group"](registration)
                buttons = registration.message.edits[0][1].inline_keyboard
                assert [row[0].callback_data for row in buttons] == [
                    *[f"subgroup:{course}:{group}:{n}" for n in numbers],
                    f"subgroup:{course}:{group}:all",
                ]
                for n in numbers:
                    await callbacks(router, "callback_query")["subgroup"](
                        FakeCallback(f"subgroup:{course}:{group}:{n}")
                    )
                    assert users.saved == (7, course, group, n)
            for number in [None, *numbers]:
                user = SimpleNamespace(
                    telegram_id=7, course=course, group_name=group, subgroup=number
                )
                users = Users(user)
                services = ScheduleService(schedule)
                router = handlers.build_router(users, services, timezone)
                delivered = []

                def capture_day(day, items, delivered=delivered):
                    delivered.extend(items)
                    return "schedule"

                def capture_schedule(items, empty, delivered=delivered):
                    delivered.extend(items)
                    return "schedule"

                def capture_ics(group, items, tz, delivered=delivered):
                    delivered.extend(items)
                    return b"calendar"

                monkeypatch.setattr(handlers, "format_day", capture_day)
                monkeypatch.setattr(handlers, "format_schedule", capture_schedule)
                monkeypatch.setattr(calendar, "build_ics", capture_ics)
                expected = [
                    x for x in lessons if x.subgroup is None or number in (None, x.subgroup)
                ]
                for day in sorted({x.date for x in lessons}):

                    class ScheduleDate(datetime):
                        @classmethod
                        def now(cls, tz=None, day=day):
                            return cls(day.year, day.month, day.day, 12, tzinfo=tz)

                    monkeypatch.setattr(handlers, "datetime", ScheduleDate)
                    delivered.clear()
                    await callbacks(router, "message")["today"](FakeMessage())
                    assert Counter(delivered) == Counter(x for x in expected if x.date == day), (
                        group,
                        number,
                        day,
                        "day",
                    )
                delivered.clear()
                weeks = services.available_weeks(group, min(x.date for x in lessons))
                for monday in weeks:
                    await callbacks(router, "callback_query")["selected_week"](
                        FakeCallback(f"week:{monday.isoformat()}")
                    )
                assert Counter(delivered) == Counter(expected), (group, number, "week")
                delivered.clear()
                service = calendar.CalendarService(users, services, timezone, None)
                await service.export_for_user(7)
                assert Counter(delivered) == Counter(expected), (group, number, "calendar")
